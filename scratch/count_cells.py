import openpyxl

wb = openpyxl.load_workbook('uploads/projects/67/Control23_Source_Logic_2.xlsx', data_only=True)
total_rows = 0
total_cells = 0
populated_cells = 0
sheet_stats = {}

for name in wb.sheetnames:
    ws = wb[name]
    s_rows = ws.max_row or 0
    s_cols = ws.max_column or 0
    s_cells = s_rows * s_cols
    s_pop = 0
    for r in range(1, s_rows + 1):
        for c in range(1, s_cols + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip() != '':
                s_pop += 1
    total_rows += s_rows
    total_cells += s_cells
    populated_cells += s_pop
    sheet_stats[name] = {'rows': s_rows, 'cols': s_cols, 'cells': s_cells, 'populated': s_pop}

print(f"Total Sheets: {len(wb.sheetnames)}")
print(f"Total Rows: {total_rows}")
print(f"Total Cells Grid: {total_cells}")
print(f"Total Populated Cells: {populated_cells}")
for s, st in sheet_stats.items():
    print(f"  {s}: Rows={st['rows']}, Cols={st['cols']}, Cells={st['cells']}, Populated={st['populated']}")
