import openpyxl

EXCEL_PATH = r"uploads/projects/67/Control23_Source_Logic_2.xlsx"
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

for sname in wb.sheetnames:
    sheet = wb[sname]
    print("\n" + "#" * 80)
    print(f"SHEET: {sname} (Rows: {sheet.max_row}, Cols: {sheet.max_column})")
    print("#" * 80)
    for r in range(1, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        # Filter trailing Nones
        while vals and vals[-1] is None:
            vals.pop()
        if vals:
            str_vals = [str(v).replace("\n", "\\n") if v is not None else "" for v in vals]
            print(f"R{r:03d}: " + " | ".join(str_vals))
