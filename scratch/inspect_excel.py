import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\Hp\Downloads\Control_Source_Logic.xlsx', data_only=True)
print("=== SHEETS IN WORKBOOK ===")
print(wb.sheetnames)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n========================================================")
    print(f"SHEET: {name} (rows: {ws.max_row}, cols: {ws.max_column})")
    print(f"========================================================")
    for r in range(1, min(10, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
        if any(row_vals):
            print(f"Row {r}: " + " | ".join(row_vals))
