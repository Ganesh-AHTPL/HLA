import openpyxl
import json
import os
import sys
import glob

if len(sys.argv) > 1:
    EXCEL_PATH = sys.argv[1]
else:
    found = glob.glob("uploads/projects/*/*.xlsx")
    EXCEL_PATH = found[0] if found else "uploads/sample.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
wb_data = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

print("=" * 80)
print("EXHAUSTIVE SCAN OF COMPLETE WORKBOOK:", EXCEL_PATH)
print("=" * 80)

total_rows_all = 0
total_cells_all = 0
total_populated_cells_all = 0

sheet_stats = {}

for sname in wb.sheetnames:
    sheet = wb[sname]
    sheet_data = wb_data[sname]
    max_r = sheet.max_row
    max_c = sheet.max_column
    pop_cells = 0
    raw_rows = []
    
    for r in range(1, max_r + 1):
        row_vals = []
        for c in range(1, max_c + 1):
            cell = sheet.cell(r, c)
            cell_data = sheet_data.cell(r, c)
            val = cell_data.value
            formula = cell.value if str(cell.value).startswith("=") else None
            
            if val is not None and str(val).strip() != "":
                pop_cells += 1
            row_vals.append({
                "col": c,
                "col_letter": openpyxl.utils.get_column_letter(c),
                "val": val,
                "formula": formula,
                "type": cell.data_type
            })
        raw_rows.append(row_vals)
    
    total_rows_all += max_r
    total_cells_all += (max_r * max_c)
    total_populated_cells_all += pop_cells
    
    sheet_stats[sname] = {
        "max_row": max_r,
        "max_col": max_c,
        "populated_cells": pop_cells,
        "rows": raw_rows
    }
    print(f"Sheet: '{sname}' -> Rows: {max_r}, Cols: {max_c}, Populated Cells: {pop_cells}")

print("\nTOTALS:")
print(f"  Sheets: {len(wb.sheetnames)}")
print(f"  Total Rows: {total_rows_all}")
print(f"  Total Cells: {total_cells_all}")
print(f"  Total Populated Cells: {total_populated_cells_all}")
print("=" * 80)
