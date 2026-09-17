import openpyxl
import re

def parse_report_derivation_logic(ws) -> list:
    """
    Parses 'Report Derivation Logic' sheet dynamically.
    Reads EVERY report section, every row and column.
    Extracts report name, attribute, source field, source table, derivation logic, remarks.
    """
    if ws is None:
        return []

    reports = []
    current_report = None
    headers = {}

    for r in range(1, (ws.max_row or 100) + 1):
        c1_val = str(ws.cell(r, 1).value or '').strip()
        c2_val = str(ws.cell(r, 2).value or '').strip()
        
        # Check if this row is a report section header
        if 'report' in c1_val.lower() and ('derivation' in c1_val.lower() or 'attribute' in c1_val.lower() or 'report' in c1_val.lower()):
            clean_name = re.sub(r'\s*-\s*attribute\s*derivation', '', c1_val, flags=re.IGNORECASE).strip()
            current_report = {
                "report_name": clean_name,
                "header_row": r,
                "attributes": []
            }
            reports.append(current_report)
            headers = {}
            continue

        # Check if this row is a table column header row (SNO, Attribute / Column Name, Source Field, etc.)
        norm_tokens = [re.sub(r'[^a-z0-9]', '', str(ws.cell(r, c).value or '').lower()) for c in range(1, (ws.max_column or 10) + 1)]
        if any('sno' in t or 'slno' in t for t in norm_tokens) and any('attribute' in t or 'column' in t for t in norm_tokens):
            headers = {}
            for c in range(1, (ws.max_column or 10) + 1):
                h_val = str(ws.cell(r, c).value or '').strip()
                h_norm = re.sub(r'[^a-z0-9]', '', h_val.lower())
                if 'sno' in h_norm or 'slno' in h_norm:
                    headers['sno'] = c
                elif 'attribute' in h_norm or 'column' in h_norm:
                    headers['attribute'] = c
                elif 'field' in h_norm:
                    headers['source_field'] = c
                elif 'table' in h_norm:
                    headers['source_table'] = c
                elif 'remark' in h_norm:
                    headers['remark'] = c
            continue

        # Check if this row is a data attribute row
        if current_report is not None and c1_val and c1_val.isdigit():
            attr_name = str(ws.cell(r, headers.get('attribute', 2)).value or '').strip()
            src_field = str(ws.cell(r, headers.get('source_field', 3)).value or '').strip()
            src_table = str(ws.cell(r, headers.get('source_table', 4)).value or '').strip()
            remark = str(ws.cell(r, headers.get('remark', 5)).value or '').strip()

            if attr_name:
                current_report["attributes"].append({
                    "row_number": r,
                    "sno": int(c1_val),
                    "attribute_name": attr_name,
                    "source_field": src_field,
                    "source_table": src_table,
                    "derivation_logic": src_field if "derive" in src_field.lower() or "count" in src_field.lower() or "case" in src_field.lower() else remark,
                    "remark": remark,
                    "kri_relationship": "Good Cases" if "good case" in remark.lower() else ("KRI Exception" if "kri" in remark.lower() or "issue" in remark.lower() else "Standard Reconciliation")
                })

    return reports


def parse_config_tables_complete(ws) -> list:
    """
    Parses 'Config Tables' sheet completely.
    Extracts every single configuration and exclusion value without truncation.
    """
    if ws is None:
        return []

    config_tables = []
    current_table = None
    current_col_name = None
    current_desc = ""
    current_rows = []

    for r in range(1, (ws.max_row or 150) + 1):
        c1 = str(ws.cell(r, 1).value or '').strip()
        if not c1:
            continue

        c1_low = c1.lower()
        if any(term in c1_low for term in ['internal profile', 'managed service', 'ip exclusion', 'customer name exclusion', 'test/dummy ip', 'internal / test customer', 'configuration & exclusion']):
            if current_table and current_rows:
                config_tables.append({
                    "table_name": current_table,
                    "config_table_name": current_table,
                    "target_column": current_col_name or "config_value",
                    "description": current_desc,
                    "record_count": len(current_rows),
                    "all_values": list(current_rows),
                    "sample_values": current_rows[:5]
                })
                current_rows = []

            if "internal profile" in c1_low:
                current_table = "cfg_internal_profiles_vdom"
                current_col_name = "hostname"
                current_desc = "Internal Profiles (VDOM) - Hostnames excluded by Rule R9"
            elif "managed service" in c1_low:
                current_table = "cfg_managed_service_types"
                current_col_name = "managed_service"
                current_desc = "Managed Service Types - Used in KRI/Non-KRI logic (B5.5)"
            elif "test/dummy ip" in c1_low or "ip exclusion" in c1_low:
                current_table = "cfg_test_dummy_ips"
                current_col_name = "ip_address"
                current_desc = "Test/Dummy IP Exclusion List - Filtered by Rule R10 (CMDB)"
            elif "customer name exclusion" in c1_low or "internal / test customer" in c1_low:
                current_table = "cfg_internal_customer_exclusions"
                current_col_name = "customer_name"
                current_desc = "Internal / Test Customer Name Exclusion List"
            else:
                current_table = "cfg_" + re.sub(r'[^a-z0-9_]', '_', c1_low)[:30].strip('_')
                current_col_name = "config_value"
                current_desc = c1
            continue

        if current_col_name and re.sub(r'[^a-z0-9]', '', c1_low) == re.sub(r'[^a-z0-9]', '', current_col_name.lower()):
            continue

        if current_table:
            current_rows.append(c1)

    if current_table and current_rows:
        config_tables.append({
            "table_name": current_table,
            "config_table_name": current_table,
            "target_column": current_col_name or "config_value",
            "description": current_desc,
            "record_count": len(current_rows),
            "all_values": list(current_rows),
            "sample_values": current_rows[:5]
        })

    return config_tables

wb = openpyxl.load_workbook('uploads/projects/67/Control23_Source_Logic_2.xlsx', data_only=True)
reps = parse_report_derivation_logic(wb['Report Derivation Logic'])
print(f"Parsed {len(reps)} report sections:")
for rep in reps:
    print(f"  - {rep['report_name']}: {len(rep['attributes'])} attributes")

cfgs = parse_config_tables_complete(wb['Config Tables'])
print(f"\nParsed {len(cfgs)} config tables:")
for cfg in cfgs:
    print(f"  - {cfg['table_name']} ({cfg['target_column']}): {cfg['record_count']} rows: {cfg['all_values'][:3]}...")
