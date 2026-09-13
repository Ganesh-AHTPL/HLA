import re
import openpyxl

def inspect_excel_structure(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    print("Sheets in workbook:", wb.sheetnames)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- Sheet: {sheet_name} (rows={ws.max_row}, cols={ws.max_column}) ---")
        for r in range(1, min(ws.max_row + 1, 15)):
            row_vals = [str(ws.cell(r, c).value).strip() if ws.cell(r, c).value is not None else "" for c in range(1, ws.max_column + 1)]
            if any(row_vals):
                print(f"Row {r:2d}: {[v for v in row_vals if v][:6]}")

if __name__ == "__main__":
    inspect_excel_structure(r"C:\Users\Hp\Downloads\Control23_Source_Logic.xlsx")
