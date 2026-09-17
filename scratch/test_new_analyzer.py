"""
Prototype script to test comprehensive 13-table parser against HLA_Solution_Design_Template.docx
"""
import os
import re
import json
import docx

TEMPLATE_PATH = r"C:\Users\Hp\Downloads\HLA_Solution_Design_Template.docx"

def clean_text(text):
    if not text:
        return ""
    return text.replace("\n", " ").strip()

def get_table_headers(table):
    if not table.rows:
        return []
    return [clean_text(c.text).lower() for c in table.rows[0].cells]

def test_parse():
    doc = docx.Document(TEMPLATE_PATH)
    print(f"Loaded document with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

    # Table 1: Version History
    version_history = []
    # Table 2: Reviewers / Approvers
    reviewers = []
    # Table 3: Control Identification
    control_id_data = {}
    # Table 4: DB Connection Details
    db_connections = []
    # Table 5: Default Control Columns
    default_columns = []
    # Table 6: Source Systems
    sources = []
    # Table 7: Data Extraction Requirements
    extractions = []
    # Table 8: Filter Buckets & Source Table Rules
    filter_rules = []
    # Table 9: Filtration Procedure
    filtration_steps = []
    # Table 10: Report Table Generation
    mappings = []
    # Table 11: Report Attributes & Numbers
    report_attributes = []
    # Table 12: Configuration Values
    config_tables = []
    # Table 13: Report Inventory & Generation Order
    report_inventory = []

    for idx, table in enumerate(doc.tables):
        headers = get_table_headers(table)
        hdr_str = " | ".join(headers)
        
        # 1. Version History
        if "version" in headers and "author" in headers:
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    version_history.append({
                        "version": cells[0] if len(cells) > 0 else "",
                        "date": cells[1] if len(cells) > 1 else "",
                        "author": cells[2] if len(cells) > 2 else "",
                        "description": cells[3] if len(cells) > 3 else ""
                    })

        # 2. Reviewers / Approvers
        elif "role" in headers and "team" in headers and ("name" in headers or "reviewer" in hdr_str):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    reviewers.append({
                        "name": cells[0] if len(cells) > 0 else "",
                        "role": cells[1] if len(cells) > 1 else "",
                        "team": cells[2] if len(cells) > 2 else ""
                    })

        # 3. Control Identification
        elif "field" in headers and "value" in headers and any("control" in clean_text(r.cells[0].text).lower() for r in table.rows[1:3]):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 2 and cells[0]:
                    k = cells[0].lower().replace(" ", "_")
                    control_id_data[k] = cells[1]

        # 4. DB Connection Details
        elif "parameter" in headers and "value" in headers and any("server" in clean_text(r.cells[0].text).lower() or "environment" in clean_text(r.cells[0].text).lower() for r in table.rows[1:3]):
            conn_entry = {}
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if len(cells) >= 2 and cells[0]:
                    param_name = cells[0].lower()
                    if "environment" in param_name:
                        conn_entry["environment"] = cells[1]
                    elif "server" in param_name or "host" in param_name:
                        conn_entry["host"] = cells[1]
                    elif "port" in param_name:
                        conn_entry["port"] = cells[1]
                    elif "database" in param_name:
                        conn_entry["database_name"] = cells[1]
                    elif "schema" in param_name:
                        conn_entry["schema_name"] = cells[1]
                    elif "credential" in param_name:
                        conn_entry["credential_reference"] = cells[1]
                    elif "connection type" in param_name:
                        conn_entry["connection_type"] = cells[1]
            if conn_entry:
                db_connections.append(conn_entry)

        # 5. Default Control Columns
        elif any("control_id" in clean_text(r.cells[0].text).upper() for r in table.rows[1:3]) or ("column name" in headers and "data type" in headers):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells) and cells[0] and not cells[0].startswith("<"):
                    default_columns.append({
                        "column_name": cells[0],
                        "data_type": cells[1] if len(cells) > 1 else "VARCHAR",
                        "description": cells[2] if len(cells) > 2 else ""
                    })

        # 6. Source Systems & Table Details (Table 6)
        elif any("source system" in h for h in headers) or (any("table name" in h for h in headers) and any("load type" in h or "type of load" in h for h in headers)):
            col_table = next((i for i, h in enumerate(headers) if "table name" in h), -1)
            col_src_sys = next((i for i, h in enumerate(headers) if "source system" in h or "system" in h), -1)
            col_db = next((i for i, h in enumerate(headers) if "db" in h or "schema" in h), -1)
            col_freq = next((i for i, h in enumerate(headers) if "refresh" in h or "frequency" in h), -1)
            col_load = next((i for i, h in enumerate(headers) if "load" in h), -1)

            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if not any(cells):
                    continue

                tbl_name = cells[col_table] if col_table >= 0 and col_table < len(cells) else cells[0]
                src_sys = cells[col_src_sys] if col_src_sys >= 0 and col_src_sys < len(cells) else ""
                db_schema = cells[col_db] if col_db >= 0 and col_db < len(cells) else ""
                freq = cells[col_freq] if col_freq >= 0 and col_freq < len(cells) else "Daily"
                load_type = cells[col_load] if col_load >= 0 and col_load < len(cells) else "Truncate & Load"

                # Parse schema vs table
                source_schema = "public"
                source_table = tbl_name
                if "." in tbl_name:
                    p = tbl_name.split(".", 1)
                    source_schema, source_table = p[0].strip(), p[1].strip()
                elif "." in db_schema:
                    p = db_schema.split(".", 1)
                    source_schema = p[1].strip()

                load_category = "Append" if "append" in load_type.lower() else "Truncate and load"

                sources.append({
                    "full_table_name": f"{source_schema}.{source_table}" if source_schema != "public" else source_table,
                    "source_system": src_sys,
                    "source_db": db_schema.split(".")[0] if "." in db_schema else db_schema,
                    "source_schema": source_schema,
                    "source_table": source_table,
                    "frequency": freq,
                    "refresh_time": freq,
                    "type_of_load": load_category,
                    "raw_load_notes": load_type
                })

        # 7. Data Extraction Requirements (Table 7)
        elif any("extraction" in h for h in headers):
            col_src = next((i for i, h in enumerate(headers) if "source" in h or "table" in h), 0)
            col_ext_type = next((i for i, h in enumerate(headers) if "extraction type" in h or "type" in h), 1)
            col_logic = next((i for i, h in enumerate(headers) if "logic" in h or "window" in h), 2)
            col_just = next((i for i, h in enumerate(headers) if "justification" in h), 3)

            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    extractions.append({
                        "source_table": cells[col_src] if col_src < len(cells) else "",
                        "extraction_type": cells[col_ext_type] if col_ext_type < len(cells) else "Full Dump",
                        "extraction_logic": cells[col_logic] if col_logic < len(cells) else "",
                        "justification": cells[col_just] if col_just < len(cells) else ""
                    })

        # 8. Filter Buckets & Source Table Rules (Table 8)
        elif any("bucket" in h or "balanced records" in h or "filter rule applied" in h for h in headers):
            col_src = next((i for i, h in enumerate(headers) if "source" in h or "table" in h), 0)
            col_rule = next((i for i, h in enumerate(headers) if "rule" in h), 1)
            col_flt = next((i for i, h in enumerate(headers) if "filtered" in h or "exception" in h), 2)
            col_bal = next((i for i, h in enumerate(headers) if "balanced" in h or "pass-through" in h), 3)

            for r_idx, row in enumerate(table.rows[1:]):
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    src_tbl = cells[col_src] if col_src < len(cells) else ""
                    rule_text = cells[col_rule] if col_rule < len(cells) else ""
                    flt_bucket = cells[col_flt] if col_flt < len(cells) else ""
                    bal_bucket = cells[col_bal] if col_bal < len(cells) else ""

                    filter_rules.append({
                        "rule_id": f"R{len(filter_rules)+1}",
                        "data_stream": src_tbl,
                        "category": "Pre-Execution Filter",
                        "rule_statement": rule_text,
                        "filtered_out_bucket": flt_bucket,
                        "balanced_bucket": bal_bucket,
                        "source_table": src_tbl
                    })

        # 9. Filtration Procedure (Table 9)
        elif "step" in headers and any("description" in h for h in headers) and any("logic" in h or "reference" in h for h in headers):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    filtration_steps.append({
                        "step": cells[0] if len(cells) > 0 else "",
                        "description": cells[1] if len(cells) > 1 else "",
                        "logic_reference": cells[2] if len(cells) > 2 else ""
                    })

        # 10. Report Table Generation (Table 10)
        elif any("attribute" in h or "column name" in h for h in headers) and any("derivation" in h or "source field" in h or "source type" in h for h in headers):
            col_tbl = next((i for i, h in enumerate(headers) if "report table" in h or "target table" in h), 0)
            col_attr = next((i for i, h in enumerate(headers) if "attribute" in h or "column" in h), 1)
            col_type = next((i for i, h in enumerate(headers) if "source type" in h or "type" in h), 2)
            col_logic = next((i for i, h in enumerate(headers) if "derivation" in h or "logic" in h or "field" in h), 3)

            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if not any(cells):
                    continue

                tbl_name = cells[col_tbl] if col_tbl < len(cells) else "target_report"
                attr_name = cells[col_attr] if col_attr < len(cells) else cells[0]
                src_type = cells[col_type] if col_type < len(cells) else "Direct"
                logic_field = cells[col_logic] if col_logic < len(cells) else ""

                is_derived = "derived" in src_type.lower() or "case" in logic_field.lower() or "if" in logic_field.lower()
                mappings.append({
                    "target_column": attr_name,
                    "report_table_name": tbl_name,
                    "source_field": logic_field if not is_derived else "(Derived Expression)",
                    "source_table": tbl_name,
                    "mapping_type": "Derived" if is_derived else "Direct",
                    "derivation_logic": logic_field if is_derived else f"Direct 1:1 mapping from {logic_field}",
                    "raw_source_type": src_type
                })

        # 11. Report Attributes & Report Numbers (Table 11)
        elif any("report name" in h for h in headers) and any("metrics" in h or "numbers" in h or "kri" in h for h in headers):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    report_attributes.append({
                        "report_name": cells[0] if len(cells) > 0 else "",
                        "key_attributes": cells[1] if len(cells) > 1 else "",
                        "metrics": cells[2] if len(cells) > 2 else ""
                    })

        # 12. Configuration Values (Table 12)
        elif any("config" in h for h in headers) and any("purpose" in h for h in headers):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    config_tables.append({
                        "config_table_name": cells[0] if len(cells) > 0 else "",
                        "purpose": cells[1] if len(cells) > 1 else "",
                        "sample_fields": cells[2] if len(cells) > 2 else "",
                        "consumed_by": cells[3] if len(cells) > 3 else ""
                    })

        # 13. Report Inventory & Generation Order (Table 13)
        elif any("sequence" in h for h in headers) and any("depends on" in h for h in headers):
            for row in table.rows[1:]:
                cells = [clean_text(c.text) for c in row.cells]
                if any(cells):
                    report_inventory.append({
                        "sequence": cells[0] if len(cells) > 0 else "",
                        "report_name": cells[1] if len(cells) > 1 else "",
                        "depends_on": cells[2] if len(cells) > 2 else "",
                        "notes": cells[3] if len(cells) > 3 else ""
                    })

    print(f"Extraction results:")
    print(f"  Table 1 Version History: {len(version_history)} entries")
    print(f"  Table 2 Reviewers: {len(reviewers)} entries")
    print(f"  Table 3 Control ID: {control_id_data}")
    print(f"  Table 4 DB Connections: {len(db_connections)} entries -> {db_connections}")
    print(f"  Table 5 Default Columns: {len(default_columns)} entries")
    print(f"  Table 6 Sources: {len(sources)} entries -> {sources}")
    print(f"  Table 7 Extractions: {len(extractions)} entries -> {extractions}")
    print(f"  Table 8 Filter Rules: {len(filter_rules)} entries -> {filter_rules}")
    print(f"  Table 9 Filtration Steps: {len(filtration_steps)} entries -> {filtration_steps}")
    print(f"  Table 10 Mappings: {len(mappings)} entries -> {mappings}")
    print(f"  Table 11 Report Attributes: {len(report_attributes)} entries -> {report_attributes}")
    print(f"  Table 12 Config Tables: {len(config_tables)} entries -> {config_tables}")
    print(f"  Table 13 Report Inventory: {len(report_inventory)} entries -> {report_inventory}")

if __name__ == "__main__":
    test_parse()
