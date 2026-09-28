"""
Test Suite for Generic Document-Driven HLA Studio Architecture.
Tests document isolation, dynamic sheet/header/row discovery, semantic inference,
absence of stale/hardcoded data, and multi-document switching.
"""

import os
import tempfile
import unittest
import openpyxl

from backend.core.generic_parser import GenericExcelParser
from backend.core.semantic_analyzer import SemanticAnalyzer
from backend.ai.context.hla_context import HLAContextExtractor
from backend.ai.context.retrieval import HLARetrievalEngine
from backend.ai.nlu.intent_parser import (
    IntentParser,
    INTENT_LIST_COMPONENTS,
    INTENT_LIST_REQUIREMENTS,
    INTENT_COMPONENT_INPUT,
    INTENT_COMPONENT_OUTPUT,
    INTENT_EXPLAIN_ARCHITECTURE,
    INTENT_LIST_DEPENDENCIES,
    INTENT_BUILD_TARGET
)
from backend.ai.nlu.entity_extractor import EntityExtractor


class TestGenericDocumentIsolation(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()

        # 1. Create HLA_A_Mapping_Rules.xlsx
        cls.hla_a_path = os.path.join(cls.temp_dir.name, "HLA_A_Mapping_Rules.xlsx")
        wb_a = openpyxl.Workbook()
        ws_m = wb_a.active
        ws_m.title = "Attribute Mappings"
        ws_m.append(["Target Column", "Source Field", "Source Table", "Mapping Type", "Derivation Logic"])
        ws_m.append(["cust_id", "c_id", "raw_customers", "Direct", "Direct 1:1 pass-through"])
        ws_m.append(["cust_name", "full_name", "raw_customers", "Direct", "Direct 1:1 pass-through"])
        ws_m.append(["is_active", "status", "raw_customers", "Derived", "CASE WHEN status = 'A' THEN TRUE ELSE FALSE END"])

        ws_r = wb_a.create_sheet(title="Business Rules")
        ws_r.append(["Rule ID", "Rule Category", "Target Dataset", "Description", "Severity"])
        ws_r.append(["BR-01", "Deduplication", "raw_customers", "Deduplicate on c_id", "MANDATORY"])
        ws_r.append(["BR-02", "Validation", "raw_customers", "Filter out NULL c_id", "CRITICAL"])
        wb_a.save(cls.hla_a_path)
        wb_a.close()

        # 2. Create HLA_B_Pipeline_Stages.xlsx
        cls.hla_b_path = os.path.join(cls.temp_dir.name, "HLA_B_Pipeline_Stages.xlsx")
        wb_b = openpyxl.Workbook()
        ws_b = wb_b.active
        ws_b.title = "Pipeline Stages"
        ws_b.append(["Stage", "Process", "Dependencies", "Input Dataset", "Output Dataset"])
        ws_b.append(["Stage 1 - Ingestion", "Read raw streaming logs", "None", "kafka_stream", "stg_raw_logs"])
        ws_b.append(["Stage 2 - Cleansing", "Remove corrupted frames", "Stage 1 - Ingestion", "stg_raw_logs", "stg_clean_logs"])
        ws_b.append(["Stage 3 - Aggregation", "Hourly window aggregation", "Stage 2 - Cleansing", "stg_clean_logs", "curated_metrics"])
        wb_b.save(cls.hla_b_path)
        wb_b.close()

        # 3. Create HLA_C_Reconciliation.xlsx
        cls.hla_c_path = os.path.join(cls.temp_dir.name, "HLA_C_Reconciliation.xlsx")
        wb_c = openpyxl.Workbook()
        ws_c1 = wb_c.active
        ws_c1.title = "Reconciliation Logic"
        ws_c1.append(["Bucket", "Description", "KRI Risk Level", "Operational Action"])
        ws_c1.append(["MATCHED_YY", "Full key parity between ledger and bank", "LOW", "Auto-reconciled"])
        ws_c1.append(["DISCREPANCY_YN", "Present in ledger but missing in bank statement", "HIGH", "Trigger audit alert"])

        ws_c2 = wb_c.create_sheet(title="Exceptions")
        ws_c2.append(["Exception ID", "Match Key", "Reason", "Remediation Action"])
        ws_c2.append(["EXC-001", "tx_99812", "Amount mismatch exceeding 0.01 tolerance", "Manual review required"])
        wb_c.save(cls.hla_c_path)
        wb_c.close()

        # 4. Create HLA_D_Architecture_Requirements.xlsx
        cls.hla_d_path = os.path.join(cls.temp_dir.name, "HLA_D_Architecture_Requirements.xlsx")
        wb_d = openpyxl.Workbook()
        ws_arch = wb_d.active
        ws_arch.title = "Architecture"
        ws_arch.append(["Component", "Responsibility", "Input", "Output", "Technology", "Interaction"])
        ws_arch.append(["Landing Zone", "Ingest raw files", "External feeds", "Raw objects", "S3 / Blob", "Batch"])
        ws_arch.append(["Validation Service", "Validate records", "Raw objects", "Validated records", "Python Service", "API"])
        ws_arch.append(["Processing Layer", "Transform & aggregate", "Validated records", "Curated tables", "Spark / SQL", "Pipeline"])
        ws_arch.append(["Quality Monitor", "Data quality checks", "Curated tables", "DQ metrics", "Great Expectations", "Event"])
        ws_arch.append(["Consumer Layer", "Serve downstream consumers", "Curated tables", "Reports / API", "PostgreSQL / REST", "Query"])

        ws_req = wb_d.create_sheet(title="Requirements")
        ws_req.append(["Requirement", "Description", "Priority", "Acceptance Criteria"])
        ws_req.append(["REQ-01", "Support real-time and batch ingest", "High", "Throughput > 10k eps"])
        ws_req.append(["REQ-02", "Automate DQ rule verification", "Critical", "Zero schema discrepancies"])
        ws_req.append(["REQ-03", "Data retention policy of 7 years", "Medium", "Audit compliance"])
        ws_req.append(["REQ-04", "Role-based access control", "High", "RBAC enforced"])
        wb_d.save(cls.hla_d_path)

        # Also save in uploads/projects/4/
        disk_path = os.path.join("uploads", "projects", "4", "HLA_D_Architecture_Requirements.xlsx")
        os.makedirs(os.path.dirname(disk_path), exist_ok=True)
        wb_d.save(disk_path)
        wb_d.close()

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def test_hla_d_architecture_requirements(self):
        """Test parsing HLA_D_Architecture_Requirements.xlsx."""
        if not os.path.exists(self.hla_d_path):
            self.skipTest(f"HLA_D file not found at {self.hla_d_path}")

        res = SemanticAnalyzer.analyze_workbook(self.hla_d_path, document_id="hla_d_test_uuid")

        self.assertEqual(res["document_id"], "hla_d_test_uuid")
        self.assertEqual(res["sheet_names"], ["Architecture", "Requirements"])
        self.assertEqual(res["detected_sections"], ["architecture_components", "requirements"])

        # Verify Components (5 components)
        comps = res["components"]
        self.assertEqual(len(comps), 5)
        comp_names = [c["component"] for c in comps]
        self.assertEqual(
            comp_names,
            ["Landing Zone", "Validation Service", "Processing Layer", "Quality Monitor", "Consumer Layer"]
        )

        val_service = next(c for c in comps if c["component"] == "Validation Service")
        self.assertEqual(val_service["input"], "Raw objects")
        self.assertEqual(val_service["output"], "Validated records")
        self.assertEqual(val_service["technology"], "Python Service")
        self.assertEqual(val_service["interaction"], "API")

        # Verify Requirements (4 requirements)
        reqs = res["requirements"]
        self.assertEqual(len(reqs), 4)
        req_ids = [r["requirement"] for r in reqs]
        self.assertEqual(req_ids, ["REQ-01", "REQ-02", "REQ-03", "REQ-04"])

        # Verify NO manufactured control
        self.assertIsNone(res["control_overview"])

        # Verify NO stale Control 23 / legacy data
        res_str = str(res)
        self.assertNotIn("dl_itsm_cmdb_daily_dump", res_str)
        self.assertNotIn("dl_pearl_active_profiles", res_str)
        self.assertNotIn("copf_id", res_str)
        self.assertNotIn("dl_ra_order_report_daily", res_str)
        self.assertNotIn("ctrl_23", res_str.lower())

    def test_hla_a_mapping_rules(self):
        """Test parsing HLA_A_Mapping_Rules.xlsx."""
        res = SemanticAnalyzer.analyze_workbook(self.hla_a_path, document_id="hla_a_test_uuid")
        self.assertIn("attribute_mappings", res["detected_sections"])
        self.assertIn("rules", res["detected_sections"])
        self.assertEqual(len(res["mappings"]), 3)
        self.assertEqual(len(res["rules"]["business_rules"]), 2)
        self.assertEqual(len(res["components"]), 0)

    def test_hla_b_pipeline_stages(self):
        """Test parsing HLA_B_Pipeline_Stages.xlsx."""
        res = SemanticAnalyzer.analyze_workbook(self.hla_b_path, document_id="hla_b_test_uuid")
        self.assertIn("pipeline_stages", res["detected_sections"])
        self.assertEqual(len(res["pipeline_stages"]), 3)
        self.assertEqual(res["pipeline_stages"][0]["stage"], "Stage 1 - Ingestion")
        self.assertEqual(res["pipeline_stages"][1]["dependencies"], "Stage 1 - Ingestion")

    def test_hla_c_reconciliation(self):
        """Test parsing HLA_C_Reconciliation.xlsx."""
        res = SemanticAnalyzer.analyze_workbook(self.hla_c_path, document_id="hla_c_test_uuid")
        self.assertIn("reconciliation", res["detected_sections"])
        self.assertTrue(len(res["reconciliation"]) >= 2)

    def test_document_switching_isolation(self):
        """Test sequential document switching ensures complete isolation."""
        doc_d1 = SemanticAnalyzer.analyze_workbook(self.hla_d_path, document_id="uuid_d1")
        doc_a = SemanticAnalyzer.analyze_workbook(self.hla_a_path, document_id="uuid_a")
        doc_b = SemanticAnalyzer.analyze_workbook(self.hla_b_path, document_id="uuid_b")
        doc_c = SemanticAnalyzer.analyze_workbook(self.hla_c_path, document_id="uuid_c")
        doc_d2 = SemanticAnalyzer.analyze_workbook(self.hla_d_path, document_id="uuid_d2")

        # Document D1
        self.assertEqual(doc_d1["document_id"], "uuid_d1")
        self.assertEqual(len(doc_d1["components"]), 5)
        self.assertEqual(len(doc_d1["mappings"]), 0)

        # Document A
        self.assertEqual(doc_a["document_id"], "uuid_a")
        self.assertEqual(len(doc_a["components"]), 0)
        self.assertEqual(len(doc_a["mappings"]), 3)

        # Document B
        self.assertEqual(doc_b["document_id"], "uuid_b")
        self.assertEqual(len(doc_b["pipeline_stages"]), 3)
        self.assertEqual(len(doc_b["components"]), 0)

        # Document C
        self.assertEqual(doc_c["document_id"], "uuid_c")
        self.assertEqual(len(doc_c["reconciliation"]), 3)

        # Document D2
        self.assertEqual(doc_d2["document_id"], "uuid_d2")
        self.assertEqual(len(doc_d2["components"]), 5)
        self.assertEqual(len(doc_d2["requirements"]), 4)

    def test_nlu_intent_parsing(self):
        """Test natural language intent classification for generic queries."""
        intent, _ = IntentParser.parse_intent("show me the components")
        self.assertEqual(intent, INTENT_LIST_COMPONENTS)

        intent, _ = IntentParser.parse_intent("what are the requirements?")
        self.assertEqual(intent, INTENT_LIST_REQUIREMENTS)

        intent, _ = IntentParser.parse_intent("show requirement")
        self.assertEqual(intent, INTENT_LIST_REQUIREMENTS)

        intent, _ = IntentParser.parse_intent("what goes into the validation service?")
        self.assertEqual(intent, INTENT_COMPONENT_INPUT)

        intent, _ = IntentParser.parse_intent("what comes out of the processing layer?")
        self.assertEqual(intent, INTENT_COMPONENT_OUTPUT)

        intent, _ = IntentParser.parse_intent("explain this architecture")
        self.assertEqual(intent, INTENT_EXPLAIN_ARCHITECTURE)

        intent, _ = IntentParser.parse_intent("what does this document contain?")
        self.assertEqual(intent, INTENT_EXPLAIN_ARCHITECTURE)

        intent, _ = IntentParser.parse_intent("make a database structure from this")
        self.assertEqual(intent, INTENT_BUILD_TARGET)

        intent, _ = IntentParser.parse_intent("what are the dependencies?")
        self.assertEqual(intent, INTENT_LIST_DEPENDENCIES)

        intent, _ = IntentParser.parse_intent("show me the input and output")
        self.assertEqual(intent, INTENT_LIST_DEPENDENCIES)

    def test_nlu_intent_parsing_generic_and_imperfect(self):
        """Test natural language intent classification for generic queries and imperfect grammar."""
        from backend.ai.nlu.intent_parser import (
            INTENT_LIST_SOURCE_TABLES,
            INTENT_RECONCILIATION_LOGIC,
            INTENT_EXPLAIN_FAILURE,
            INTENT_SUMMARIZE_PROJECT
        )
        
        # Generic variations for source tables
        self.assertEqual(IntentParser.parse_intent("show source tables")[0], INTENT_LIST_SOURCE_TABLES)
        self.assertEqual(IntentParser.parse_intent("what are the source tables")[0], INTENT_LIST_SOURCE_TABLES)
        self.assertEqual(IntentParser.parse_intent("which tables are sources")[0], INTENT_LIST_SOURCE_TABLES)
        self.assertEqual(IntentParser.parse_intent("what tables are coming in")[0], INTENT_LIST_SOURCE_TABLES)
        self.assertEqual(IntentParser.parse_intent("show me input tables")[0], INTENT_LIST_SOURCE_TABLES)
        self.assertEqual(IntentParser.parse_intent("what are the input datasets?")[0], INTENT_LIST_SOURCE_TABLES)

        # Rules and reconciliation
        self.assertEqual(IntentParser.parse_intent("explain the rules")[0], INTENT_RECONCILIATION_LOGIC)
        self.assertEqual(IntentParser.parse_intent("how does reconciliation work?")[0], INTENT_RECONCILIATION_LOGIC)
        self.assertEqual(IntentParser.parse_intent("what happens after filtering?")[0], INTENT_RECONCILIATION_LOGIC)

        # Mappings and transformations
        self.assertEqual(IntentParser.parse_intent("show mappings")[0], INTENT_BUILD_TARGET)
        self.assertEqual(IntentParser.parse_intent("explain this transformation")[0], INTENT_BUILD_TARGET)

        # Failure & Summary
        self.assertEqual(IntentParser.parse_intent("why did this process fail?")[0], INTENT_EXPLAIN_FAILURE)
        self.assertEqual(IntentParser.parse_intent("give me a summary")[0], INTENT_SUMMARIZE_PROJECT)

        # Grammatical and informal typo variations
        self.assertEqual(IntentParser.parse_intent("wht are the tbls")[0], INTENT_LIST_SOURCE_TABLES)
        self.assertEqual(IntentParser.parse_intent("shw comps plz")[0], INTENT_LIST_COMPONENTS)
        self.assertEqual(IntentParser.parse_intent("shw me reqs")[0], INTENT_LIST_REQUIREMENTS)
        self.assertEqual(IntentParser.parse_intent("what are the deps?")[0], INTENT_LIST_DEPENDENCIES)
        self.assertEqual(IntentParser.parse_intent("how does recon work")[0], INTENT_RECONCILIATION_LOGIC)

    def test_extracted_facts_provenance(self):
        """Test that all extracted facts contain provenance back to the uploaded Excel file."""
        res_d = SemanticAnalyzer.analyze_workbook(self.hla_d_path, document_id="hla_d_prov_test")
        comps = res_d["components"]
        self.assertTrue(len(comps) > 0)
        for c in comps:
            self.assertIn("provenance", c)
            self.assertEqual(c["provenance"]["document_id"], "hla_d_prov_test")
            self.assertEqual(c["provenance"]["sheet"], "Architecture")
            self.assertIsNotNone(c["provenance"]["row"])

        reqs = res_d["requirements"]
        self.assertTrue(len(reqs) > 0)
        for r in reqs:
            self.assertIn("provenance", r)
            self.assertEqual(r["provenance"]["sheet"], "Requirements")
            self.assertIsNotNone(r["provenance"]["row"])

        res_a = SemanticAnalyzer.analyze_workbook(self.hla_a_path, document_id="hla_a_prov_test")
        maps = res_a["mappings"]
        for m in maps:
            self.assertIn("provenance", m)
            self.assertEqual(m["provenance"]["sheet"], "Attribute Mappings")


if __name__ == "__main__":
    unittest.main()
