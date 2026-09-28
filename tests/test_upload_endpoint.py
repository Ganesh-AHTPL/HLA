import io
import os
import unittest
import openpyxl
from app import app, db
from models import User, Project, Document
from auth import generate_token

class TestUploadEndpoint(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config["TESTING"] = True
        cls.client = app.test_client()

        with app.app_context():
            # Get or create admin user
            admin = User.query.filter_by(username="admin").first()
            if not admin:
                admin = User(username="admin", role="admin")
                admin.set_password("admin123")
                db.session.add(admin)
                db.session.commit()
            cls.admin_token = generate_token(admin)

            # Get or create project 4
            project = db.session.get(Project, 4)
            if not project:
                project = Project(id=4, name="New format-1", description="Test Project")
                db.session.add(project)
                db.session.commit()

    def test_upload_excel_document_project_4(self):
        # Create an in-memory workbook representing HLA_D
        wb = openpyxl.Workbook()
        ws_arch = wb.active
        ws_arch.title = "Architecture"
        ws_arch.append(["Component", "Responsibility", "Input", "Output", "Technology", "Interaction"])
        ws_arch.append(["Data Pipeline", "Ingest telemetry", "Raw Events", "Processed Logs", "Spark", "Batch"])
        ws_arch.append(["Validation Service", "Verify schema rules", "Processed Logs", "Validated Metrics", "Python", "Sync"])

        ws_req = wb.create_sheet(title="Requirements")
        ws_req.append(["Requirement", "Description", "Priority", "Acceptance Criteria"])
        ws_req.append(["REQ-01", "Support streaming ingest", "High", "Throughput > 10k eps"])

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)

        response = self.client.post(
            "/api/projects/4/upload",
            headers={"Authorization": f"Bearer {self.admin_token}"},
            data={
                "file": (excel_bytes, "HLA_Test_Doc.xlsx")
            },
            content_type="multipart/form-data"
        )

        self.assertEqual(response.status_code, 200, f"Upload failed with response: {response.get_json()}")
        data = response.get_json()
        self.assertIn("document_id", data)
        self.assertIn("analysis", data)
        self.assertIn("id", data)
        self.assertEqual(data["filename"], "HLA_Test_Doc.xlsx")

        analysis = data["analysis"]
        self.assertIn("components", analysis)
        self.assertEqual(len(analysis["components"]), 2)
        self.assertEqual(len(analysis["requirements"]), 1)

    def test_upload_existing_hla_d_file(self):
        wb = openpyxl.Workbook()
        ws_arch = wb.active
        ws_arch.title = "Architecture"
        ws_arch.append(["Component", "Responsibility", "Input", "Output", "Technology", "Interaction"])
        ws_arch.append(["Landing Zone", "Ingest raw files", "External feeds", "Raw objects", "S3 / Blob", "Batch"])
        ws_arch.append(["Validation Service", "Validate records", "Raw objects", "Validated records", "Python Service", "API"])
        ws_arch.append(["Processing Layer", "Transform & aggregate", "Validated records", "Curated tables", "Spark / SQL", "Pipeline"])
        ws_arch.append(["Quality Monitor", "Data quality checks", "Curated tables", "DQ metrics", "Great Expectations", "Event"])
        ws_arch.append(["Consumer Layer", "Serve downstream consumers", "Curated tables", "Reports / API", "PostgreSQL / REST", "Query"])

        ws_req = wb.create_sheet(title="Requirements")
        ws_req.append(["Requirement", "Description", "Priority", "Acceptance Criteria"])
        ws_req.append(["REQ-01", "Support real-time and batch ingest", "High", "Throughput > 10k eps"])
        ws_req.append(["REQ-02", "Automate DQ rule verification", "Critical", "Zero schema discrepancies"])
        ws_req.append(["REQ-03", "Data retention policy of 7 years", "Medium", "Audit compliance"])
        ws_req.append(["REQ-04", "Role-based access control", "High", "RBAC enforced"])

        excel_bytes = io.BytesIO()
        wb.save(excel_bytes)
        excel_bytes.seek(0)

        # Also ensure uploads/projects/4/HLA_D_Architecture_Requirements.xlsx is saved correctly on disk
        disk_path = os.path.join("uploads", "projects", "4", "HLA_D_Architecture_Requirements.xlsx")
        os.makedirs(os.path.dirname(disk_path), exist_ok=True)
        wb.save(disk_path)

        response = self.client.post(
            "/api/projects/4/upload",
            headers={"Authorization": f"Bearer {self.admin_token}"},
            data={
                "file": (excel_bytes, "HLA_D_Architecture_Requirements.xlsx")
            },
            content_type="multipart/form-data"
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        self.assertIn("document_id", data)
        self.assertIn("analysis", data)
        analysis = data["analysis"]
        self.assertEqual(len(analysis.get("components", [])), 5)
        self.assertEqual(len(analysis.get("requirements", [])), 4)
        self.assertEqual(data["analysis"]["document_id"], data["document_id"])

if __name__ == "__main__":
    unittest.main()
